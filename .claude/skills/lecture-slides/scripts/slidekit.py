"""python-pptx 슬라이드 헬퍼 (16:9, 맑은 고딕, 남색 테마)."""
import json, os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

FONT = "맑은 고딕"


def rgb(h):
    h = h.lstrip("#"); return RGBColor(int(h[:2], 16), int(h[2:4], 16), int(h[4:], 16))


class Deck:
    NAVY = "#1F3A5F"; ACCENT = "#E4572E"; TEAL = "#2A9D8F"; GOLD = "#F4A261"; GRAY = "#8D99AE"; LIGHT = "#EEF2F7"
    COLORS = [ACCENT, GOLD, NAVY, TEAL, "#9B5DE5", "#577590"]

    def __init__(self, root, week, course_label, anim_dir=None, fig_dir=None):
        self.root = root
        wk = os.path.join(root, "lectures", f"week{week:02d}")
        self.anim = anim_dir or os.path.join(wk, "anim")
        self.fig = fig_dir or os.path.join(wk, "fig")
        os.makedirs(self.fig, exist_ok=True)
        self.img = os.path.join(root, "01_source", "이미지_도표")
        cj = os.path.join(self.img, "credits.json")
        self.credits = json.load(open(cj, encoding="utf-8")) if os.path.exists(cj) else {}
        self.label = course_label
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.n = 0

    # ---- 기본 요소 ----
    def text(self, s, x, y, w, h, lines, size=20, color=None, bold=False):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line; p.space_after = Pt(10)
            for r in p.runs:
                r.font.size = Pt(size); r.font.name = FONT; r.font.bold = bold; r.font.color.rgb = rgb(color or self.NAVY)
        return tb

    def bullets(self, s, x, y, w, h, items, size=20):
        return self.text(s, x, y, w, h, ["•  " + it for it in items], size=size)

    def pic(self, s, path, x, y, w=None, h=None):
        return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)

    def box(self, s, x, y, w, h, title, body, color=None, size=24):
        color = color or self.ACCENT
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(self.LIGHT); sh.line.color.rgb = rgb(color); sh.line.width = Pt(2)
        sh.adjustments[0] = 0.08
        self.text(s, x + 0.25, y + 0.15, w - 0.5, 0.5, title, size=16, color=color, bold=True)
        self.text(s, x + 0.25, y + 0.65, w - 0.5, h - 0.8, body, size=size, bold=True)

    def notes(self, s, txt):
        s.notes_slide.notes_text_frame.text = txt

    # ---- 슬라이드 ----
    def _new(self):
        self.n += 1
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def cover(self, kicker, title, subtitle, photo=None):
        s = self._new()
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height)
        bg.fill.solid(); bg.fill.fore_color.rgb = rgb(self.NAVY); bg.line.fill.background()
        self.text(s, 0.9, 2.2, 6, 0.6, kicker, size=20, color=self.GOLD, bold=True)
        self.text(s, 0.9, 2.9, 6.5, 1.2, title, size=48, color="#FFFFFF", bold=True)
        self.text(s, 0.9, 4.1, 6, 0.6, subtitle, size=24, color="#DDE3EA")
        if photo:
            self.pic(s, self.crop(photo, 5.6, 7.5), 7.733, 0, 5.6, 7.5)
        return s

    def slide(self, title, tag=None):
        s = self._new()
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = rgb(self.NAVY); bar.line.fill.background()
        if tag:
            self.text(s, 0.6, 0.35, 8, 0.4, tag, size=13, color=self.ACCENT, bold=True)
        self.text(s, 0.6, 0.65, 12, 0.9, title, size=30, bold=True)
        self.text(s, 0.6, 7.0, 8, 0.3, self.label, size=10, color=self.GRAY)
        self.text(s, 12.0, 7.0, 0.8, 0.3, str(self.n), size=10, color=self.GRAY)
        return s

    # ---- 사진, 영상 ----
    def crop(self, name, w, h):
        """사진을 칸 비율(w:h)로 가운데 자르기"""
        im = Image.open(os.path.join(self.img, name)).convert("RGB")
        W, H = im.size; target = w / h
        if W / H > target:
            nw = int(H * target); bx = ((W - nw) // 2, 0, (W - nw) // 2 + nw, H)
        else:
            nh = int(W / target); bx = (0, (H - nh) // 2, W, (H - nh) // 2 + nh)
        p = os.path.join(self.fig, f"crop_{w:.1f}x{h:.1f}_{name}")
        im.crop(bx).save(p, quality=90)
        return p

    def photo(self, s, name, x, y, w, h):
        """비유 사진 + 하단 출처"""
        self.pic(s, self.crop(name, w, h), x, y, w, h)
        self.text(s, x, y + h + 0.02, w, 0.3, self.credits.get(name, "출처 미기록"), size=9, color=self.GRAY)

    def analogy_tag(self, s, label="일상 속 비유"):
        t = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.9), Inches(0.4), Inches(1.8), Inches(0.5))
        t.fill.solid(); t.fill.fore_color.rgb = rgb(self.GOLD); t.line.fill.background()
        tf = t.text_frame; tf.text = label
        r = tf.paragraphs[0].runs[0]; r.font.size = Pt(14); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = rgb("#FFFFFF")

    def movie(self, s, scene, x, y, w):
        """manim 영상 (클릭 재생, 포스터 = 마지막 장면)"""
        h = w * 9 / 16
        s.shapes.add_movie(os.path.join(self.anim, scene + ".mp4"), Inches(x), Inches(y), Inches(w), Inches(h),
                           poster_frame_image=os.path.join(self.anim, scene + "_poster.png"), mime_type="video/mp4")
        self.text(s, x, y + h + 0.02, w, 0.3, "▶ 클릭하여 애니메이션 재생", size=10, color=self.GRAY)

    def save(self, path):
        path = path if os.path.isabs(path) else os.path.join(self.root, path)
        self.prs.save(path)
        print("saved", path, self.n, "slides")
